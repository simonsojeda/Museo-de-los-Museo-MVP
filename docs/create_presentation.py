from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
P=Path('/mnt/data/MuseoDeLosMuseos')
imgs=P/'app/src/main/res/drawable'
out=Path('/mnt/data/MuseoDeLosMuseos_Presentacion_DSY1105.pptx')
prs=Presentation();prs.slide_width=Inches(13.33);prs.slide_height=Inches(7.5)
NAVY=RGBColor(5,43,62);GOLD=RGBColor(189,151,88);WHITE=RGBColor(255,255,255);INK=RGBColor(20,50,68);PALE=RGBColor(241,246,248);TEAL=RGBColor(29,124,125)
def textbox(s,x,y,w,h,text,size=18,color=INK,bold=False,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color;p.space_after=Pt(6)
  if align:p.alignment=align
 return sh
def bg(s,color=WHITE):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height);sh.fill.solid();sh.fill.fore_color.rgb=color;sh.line.fill.background()
def rect(s,x,y,w,h,fill,round=True):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y), Inches(w), Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.fill.background();return sh
def title(s,head,sub=''):
 rect(s,0,0,13.33,.14,GOLD,False);textbox(s,.65,.37,12,.55,head,27,NAVY,True)
 if sub:textbox(s,.65,1,12,.38,sub,14,TEAL)
def picture(s,name,x,y,w,h):
 file=imgs/f'cover_{name}.png';s.shapes.add_picture(str(file), Inches(x), Inches(y),width=Inches(w),height=Inches(h))
def footer(s,index):textbox(s,11.7,7.1,1.1,.2,f'{index} / 9',10,TEAL)
def slide():return prs.slides.add_slide(prs.slide_layouts[6])
# 1
s=slide();bg(s,NAVY);picture(s,'museum',6.8,0,6.53,7.5);rect(s,0,0,7.15,7.5,NAVY,False)
textbox(s,.65,1.45,6,.5,'DUOC UC  ·  DSY1105',16,GOLD,True)
textbox(s,.65,2.2,6,1.65,'MUSEO DE LOS\nMUSEOS',39,WHITE,True)
textbox(s,.65,4.4,5.8,.95,'Aplicación Android en Kotlin\nMVP académico · Evaluación Parcial 2',20,WHITE)
textbox(s,.65,6.75,5,.25,'Proyecto y planificación de 5 semanas',12,GOLD)
# 2
s=slide();bg(s,PALE);title(s,'Problema y propuesta','Basado en el caso Museo de los Museos · Duoc UC')
for x,label,desc in [(0.75,'ANTES','Cursos, productos, museos virtuales y cotizaciones distribuidos en canales separados.'),(6.85,'DESPUÉS','Una aplicación unificada con catálogo, mapa, fichas, favoritos, solicitudes y envío simulado.')]:
 rect(s,x,1.7,5.7,4.55,WHITE);rect(s,x,1.7,5.7,.65,NAVY);textbox(s,x+.3,1.84,4.9,.35,label,18,WHITE,True)
 textbox(s,x+.38,2.65,4.95,2.5,desc,23,INK)
textbox(s,.85,6.6,11.7,.4,'Restricción: ningún pago, envío ni dato personal real. Uso exclusivamente académico.',15,TEAL,True);footer(s,2)
# 3
s=slide();bg(s,PALE);title(s,'15 pantallas organizadas en 5 módulos')
blocks=[('01 · Acceso','Bienvenida · Registro\nLogin · Dashboard'),('02 · Investigación','Mapa · Museo virtual\nEscáner QR · Resultado QR'),('03 · Comercio','Catálogo · Detalle\nCarrito y cálculo de envío'),('04 · Seguimiento','Favoritos · Solicitudes\nNotificaciones'),('05 · Usuario','Perfil y configuración\nPermisos y preferencias')]
for i,(label,body) in enumerate(blocks):
 row=i//3;col=i%3;x=.7+col*4.12;y=1.62+row*2.53
 rect(s,x,y,3.78,2.15,WHITE);rect(s,x,y,3.78,.53,GOLD);textbox(s,x+.2,y+.1,3.4,.3,label,16,NAVY,True);textbox(s,x+.22,y+.72,3.35,1.14,body,17,INK)
footer(s,3)
# 4
s=slide();bg(s,PALE);title(s,'Arquitectura móvil','UI, dominio y datos desacoplados para la rúbrica de MVVM')
for x,y,w,t,detail in [(0.8,2.1,3.25,'Jetpack Compose','15 pantallas y estados'),(5.02,2.1,3.25,'ViewModel','Sesión, favoritos, carrito'),(9.24,2.1,3.25,'Room / SQLite','Usuarios y registros locales'),(5.02,4.65,3.25,'Reglas puras','Validación + cotización'),(9.24,4.65,3.25,'Retrofit · REST','API sintética opcional')]:
 rect(s,x,y,w,1.43,WHITE);rect(s,x,y,w,.52,NAVY);textbox(s,x+.18,y+.13,w-.36,.3,t,16,WHITE,True);textbox(s,x+.15,y+.71,w-.3,.56,detail,15,INK)
textbox(s,4.22,2.52,.65,.4,'→',27,GOLD,True);textbox(s,8.48,2.52,.65,.4,'→',27,GOLD,True);textbox(s,6.48,3.75,.65,.4,'↓',27,GOLD,True);textbox(s,10.64,3.75,.65,.4,'↓',27,GOLD,True)
footer(s,4)
# 5
s=slide();bg(s,PALE);title(s,'Perfiles y permisos nativos')
for i,(label,desc) in enumerate([('CLIENTE','Explora · guarda favoritos\ncompra simulada · revisa historial'),('INVESTIGADOR','Crea y edita fichas de museos\nvalida referencias QR'),('ADMINISTRADOR','Gestiona catálogo, solicitudes\ny emite avisos de laboratorio')]):
 x=.65+i*4.24;rect(s,x,1.62,3.85,2.5,WHITE);rect(s,x,1.62,3.85,.6,NAVY);textbox(s,x+.15,1.78,3.55,.25,label,16,WHITE,True);textbox(s,x+.22,2.45,3.3,1.2,desc,18,INK)
rect(s,.75,4.75,5.8,1.7,RGBColor(231,240,242));textbox(s,1,4.98,5,.3,'RECURSO 1  ·  CÁMARA',19,NAVY,True);textbox(s,1,5.55,5.1,.58,'Escaneo QR con permiso y validación local',17,INK)
rect(s,6.75,4.75,5.8,1.7,RGBColor(231,240,242));textbox(s,7,4.98,5.2,.3,'RECURSO 2  ·  UBICACIÓN',19,NAVY,True);textbox(s,7,5.55,5.2,.62,'Consulta de ubicación aproximada del teléfono',17,INK);footer(s,5)
# 6
s=slide();bg(s,PALE);title(s,'Recursos gráficos incluidos en la aplicación','Ilustraciones originales creadas para demostración, no capturas de ejecución')
for i,(img,label) in enumerate([('museum','Museos'),('book','Publicaciones'),('kit','Conservación'),('course','Cursos'),('merch','Merchandising')]):
 x=.6+i*2.56;picture(s,img,x,2.05,2.30,2.35);textbox(s,x,4.65,2.3,.45,label,16,NAVY,True,PP_ALIGN.CENTER)
textbox(s,.75,5.75,12,.6,'Todos los nombres, precios, ubicaciones y recorridos son ficticios.',19,TEAL,True,PP_ALIGN.CENTER);footer(s,6)
# 7
s=slide();bg(s,PALE);title(s,'Planificación Trello: 5 semanas','36 tickets con criterios de aceptación y responsables A / B')
labels=[('S1','Base, GitHub, Room\ny sistema visual'),('S2','Registro, roles,\nvalidaciones y perfil'),('S3','Museos, mapa, QR,\nGPS y favoritos'),('S4','Catálogo, carrito,\nenvíos, REST y avisos'),('S5','Pruebas, evidencia,\nREADME y entrega')]
for i,(number,body) in enumerate(labels):
 x=.44+i*2.61
 rect(s,x,2.1,2.4,3.4,WHITE);rect(s,x,2.1,2.4,.84,NAVY);textbox(s,x+.25,2.26,1.7,.52,number,25,WHITE,True)
 textbox(s,x+.22,3.3,2.06,1.75,body,17,INK)
textbox(s,.75,6.25,12,.5,'A y B son etiquetas de asignación sugeridas; deben reemplazarse por los dos integrantes reales.',15,TEAL)
footer(s,7)
# 8
s=slide();bg(s,PALE);title(s,'Ponderación de la evaluación','Distribución indicada en la rúbrica oficial de la Evaluación Parcial 2')
metrics=[('Diseño y navegación',15),('Formularios visuales',15),('Lógica desacoplada',10),('Animaciones',10),('Modularidad + SQLite',15),('GitHub + Trello',20),('Recursos nativos',15)]
for i,(n,pct) in enumerate(metrics):
 y=1.48+i*.69
 textbox(s,.76,y,3.7,.4,n,17,INK)
 rect(s,4.6,y+.05,6.7,.23,RGBColor(212,225,231));rect(s,4.6,y+.05,6.7*pct/20,.23,GOLD)
 textbox(s,11.55,y,1.1,.38,str(pct)+'%',18,NAVY,True)
footer(s,8)
# 9
s=slide();bg(s,NAVY);textbox(s,.76,.65,10,.6,'Preparación para la defensa',30,WHITE,True)
texts=[('1','Abrir el proyecto en Android Studio y sincronizar dependencias.'),('2','Ejecutar en emulador y comprobar las 15 pantallas.'),('3','Probar cámara QR y permiso de ubicación en el equipo.'),('4','Publicar GitHub y registrar actividad real en Trello.'),('5','Subir código a AVA y cumplir la entrega individual/grupal.')]
for i,(n,t) in enumerate(texts):
 y=1.75+i*.93;rect(s,.85,y,.53,.53,GOLD);textbox(s,.98,y+.08,.24,.3,n,19,NAVY,True)
 textbox(s,1.67,y+.06,10.75,.65,t,19,WHITE)
textbox(s,.78,6.86,12,.25,'La compilación Android y la publicación real en Trello requieren verificación y conexión posterior.',12,GOLD)
prs.save(out)
print(out,out.stat().st_size)
